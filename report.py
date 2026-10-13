"""
Rapport de comparaison de portefeuilles au format .xlsx, avec de vrais graphiques natifs,
des mises en forme conditionnelles et des tableaux utiles (rendements annuels, mensuels, drawdowns…).
"""
import io
import json
import os
from datetime import datetime

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference, ScatterChart, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import backend
from i18n import T, lang

ROOT = os.path.dirname(os.path.abspath(__file__))
GREEN, GREEN_SOFT, RED_SOFT, GREY = "1E7145", "D9EAD3", "F4CCCC", "F2F3F2"
SERIES = ["1E7145", "2E75B6", "C55A11", "7F6000", "7030A0", "7F7F7F"]
THIN = Side(style="thin", color="D9DCD9")
HEAD = dict(font=Font(bold=True, color="FFFFFF"), fill=PatternFill("solid", fgColor=GREEN), alignment=Alignment(horizontal="center", vertical="center", wrap_text=True))

_names = None


def asset_name(tk):
    global _names
    if _names is None:
        try:
            _names = {a["ticker"]: a["name"] for a in json.load(open(os.path.join(ROOT, "universe.json"), encoding="utf-8"))}
        except Exception:
            _names = {}
    return _names.get(tk, tk)


def _head(ws, row, col, text, width=None):
    c = ws.cell(row=row, column=col, value=text)
    c.font, c.fill, c.alignment = HEAD["font"], HEAD["fill"], HEAD["alignment"]
    c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    if width:
        ws.column_dimensions[get_column_letter(col)].width = width
    return c


def _cell(ws, row, col, value, fmt=None, bold=False):
    c = ws.cell(row=row, column=col, value=None if value is None or (isinstance(value, float) and not np.isfinite(value)) else value)
    if fmt:
        c.number_format = fmt
    if bold:
        c.font = Font(bold=True)
    c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    return c


def _style_chart(ch, title, y_fmt=None, height=7.5, width=17):
    ch.title = title
    ch.height, ch.width = height, width
    ch.legend.position = "b"
    if y_fmt:
        ch.y_axis.numFmt = y_fmt
    ch.y_axis.delete = False
    ch.x_axis.delete = False


def build_comparison(portfolios, fetch, period="5y", start=None, end=None, benchmark="^GSPC", capital=10000.0,
                     benchmark_label=None, currency=None, risk_free_source=None):
    """portfolios : [{name, items: [[ticker, poids %], …], rebalance}]. Renvoie les octets du fichier .xlsx.
    `fetch` renvoie des cours déjà convertis dans `currency` (None = chaque actif dans sa devise de cotation)."""
    bench_name = benchmark_label or benchmark
    money = f" {currency}" if currency else ""
    all_tickers = sorted({tk for p in portfolios for tk, _ in p["items"]} | {benchmark})
    data = fetch(all_tickers, period=period, start=start, end=end)

    # 1) Valeur de chaque portefeuille, puis période commune à tous
    series = {}
    for p in portfolios:
        tks = [tk for tk, _ in p["items"]]
        w = np.array([float(x) for _, x in p["items"]]); w = w / w.sum()
        prices = pd.concat({tk: data[tk]["Close"] for tk in tks}, axis=1, join="inner").dropna()
        series[p["name"]] = backend.portfolio_value_series(prices, w, p.get("rebalance") or "none")
    values = pd.concat(series, axis=1, join="inner").dropna()
    bench = data[benchmark]["Close"].reindex(values.index).ffill().bfill()
    values[T(f"Indice de référence ({bench_name})", f"Benchmark ({bench_name})")] = bench
    if len(values) < 20:
        raise ValueError(T("Historique commun trop court pour comparer ces portefeuilles.", "Shared history too short to compare these portfolios."))
    values = values / values.iloc[0] * capital  # même capital de départ, même date
    names = list(values.columns)
    port_names = names[:-1]
    ppy = 252

    # 2) Indicateurs
    rets = values.pct_change().dropna()
    rows = []  # (libellé, format, sens « meilleur », valeurs)

    def add(label, fmt, better, fn):
        rows.append((label, fmt, better, [fn(n) for n in names]))

    metrics = {n: backend.compute_core_metrics(values[n], bench, 0.95, ppy) for n in names}
    yearly = values.resample("YE").last()
    yearly = pd.concat([values.iloc[[0]], yearly]).pct_change().dropna()
    yearly.index = yearly.index.year
    monthly = values.resample("ME").last()
    monthly = pd.concat([values.iloc[[0]], monthly]).pct_change().dropna()
    dd = values / values.cummax() - 1

    add(T(f"Valeur finale de {capital:,.0f}{money}".replace(",", " "), f"Final value of {capital:,.0f}{money}"), "#,##0", 1, lambda n: float(values[n].iloc[-1]))
    add(T("Gain / perte", "Gain / loss"), "#,##0;[Red]-#,##0", 1, lambda n: float(values[n].iloc[-1] - capital))
    add(T("Rendement total", "Total return"), "0.0%", 1, lambda n: float(values[n].iloc[-1] / capital - 1))
    add(T("Rendement annuel composé (CAGR)", "Compound annual return (CAGR)"), "0.00%", 1, lambda n: metrics[n]["cagr_pct"] / 100)
    add(T("Volatilité annualisée", "Annualised volatility"), "0.00%", -1, lambda n: metrics[n]["annualized_volatility_pct"] / 100)
    add(T("Pire baisse (max drawdown)", "Worst drop (max drawdown)"), "0.00%", 1, lambda n: float(dd[n].min()))
    add(T("Plus longue période sous un sommet (séances)", "Longest time below a peak (sessions)"), "0", -1, lambda n: metrics[n]["longest_drawdown_days"])
    add(T("Ratio de Sharpe", "Sharpe ratio"), "0.00", 1, lambda n: metrics[n]["sharpe_ratio"])
    add(T("Ratio de Sortino", "Sortino ratio"), "0.00", 1, lambda n: metrics[n]["sortino_ratio"])
    add(T("Ratio de Calmar", "Calmar ratio"), "0.00", 1, lambda n: metrics[n]["calmar_ratio"])
    add(T("Perte journalière à 95 % (VaR)", "Daily loss at 95% (VaR)"), "0.00%", 1, lambda n: metrics[n]["var_historical_pct"] / 100)
    add(T("Perte moyenne des pires jours (CVaR)", "Average loss on worst days (CVaR)"), "0.00%", 1, lambda n: metrics[n]["cvar_pct"] / 100)
    add(T("Meilleure année", "Best year"), "0.0%", 1, lambda n: float(yearly[n].max()))
    add(T("Pire année", "Worst year"), "0.0%", 1, lambda n: float(yearly[n].min()))
    add(T("Meilleur mois", "Best month"), "0.0%", 1, lambda n: float(monthly[n].max()))
    add(T("Pire mois", "Worst month"), "0.0%", 1, lambda n: float(monthly[n].min()))
    add(T("Mois positifs", "Positive months"), "0%", 1, lambda n: float((monthly[n] > 0).mean()))
    add(T("Bêta vs indice de référence", "Beta vs benchmark"), "0.00", 0, lambda n: metrics[n]["beta"])
    add(T("Corrélation à l'indice de référence", "Correlation with benchmark"), "0.00", 0, lambda n: metrics[n]["correlation"])

    wb = Workbook()
    ws = wb.active
    ws.title = T("Tableau de bord", "Dashboard")
    ws.sheet_view.showGridLines = False
    ws["A1"] = T("Comparaison de portefeuilles", "Portfolio comparison")
    ws["A1"].font = Font(bold=True, size=18, color=GREEN)
    rf_src = f" ({risk_free_source})" if risk_free_source else ""
    ccy_note = (T(f" · tous les cours sont convertis en {currency} au taux de change de chaque jour",
                  f" · all prices are converted to {currency} at each day's exchange rate") if currency
                else T(" · chaque actif reste dans sa devise de cotation (pas de conversion de change)",
                       " · each asset stays in its quote currency (no currency conversion)"))
    ws["A2"] = T(
        f"Du {values.index[0]:%d/%m/%Y} au {values.index[-1]:%d/%m/%Y} · indice de référence : {bench_name} · capital de départ {capital:,.0f}{money} · taux sans risque {backend.risk_free() * 100:.2f} %{rf_src} · exporté le {datetime.now():%d/%m/%Y %H:%M}".replace(",", " ") + ccy_note,
        f"From {values.index[0]:%Y-%m-%d} to {values.index[-1]:%Y-%m-%d} · benchmark: {bench_name} · starting capital {capital:,.0f}{money} · risk-free rate {backend.risk_free() * 100:.2f}%{rf_src} · exported {datetime.now():%Y-%m-%d %H:%M}" + ccy_note)
    ws["A2"].font = Font(italic=True, color="69716C")
    ws["A3"] = T("En vert : le meilleur portefeuille de chaque ligne · en rouge : le moins bon (indice de référence exclu).",
                 "Green: best portfolio on each row · red: weakest (benchmark excluded).")
    ws["A3"].font = Font(size=9, color="69716C")

    r0 = 5
    _head(ws, r0, 1, T("Indicateur", "Metric"), 40)
    for j, n in enumerate(names):
        _head(ws, r0, 2 + j, n, 16)
    for i, (label, fmt, better, vals) in enumerate(rows):
        r = r0 + 1 + i
        _cell(ws, r, 1, label, bold=True).fill = PatternFill("solid", fgColor=GREY)
        for j, v in enumerate(vals):
            c = _cell(ws, r, 2 + j, v, fmt)
            if j == len(names) - 1:
                c.font = Font(italic=True, color="69716C")
        pv = [v for v in vals[:-1] if v is not None and np.isfinite(v)]
        if better and len(pv) >= 2:
            best, worst = (max(pv), min(pv)) if better > 0 else (min(pv), max(pv))
            for j, v in enumerate(vals[:-1]):
                if v == best:
                    ws.cell(row=r, column=2 + j).fill = PatternFill("solid", fgColor=GREEN_SOFT)
                    ws.cell(row=r, column=2 + j).font = Font(bold=True, color=GREEN)
                elif v == worst:
                    ws.cell(row=r, column=2 + j).fill = PatternFill("solid", fgColor=RED_SOFT)
    ws.freeze_panes = "B6"

    # 3) Feuille de données (hebdomadaire pour des graphiques lisibles)
    wd = wb.create_sheet(T("Données", "Data"))
    weekly = values.resample("W-FRI").last().dropna()
    wdd = (weekly / weekly.cummax() - 1)
    _head(wd, 1, 1, "Date", 12)
    for j, n in enumerate(names):
        _head(wd, 1, 2 + j, n, 16)
        _head(wd, 1, 2 + len(names) + j, f"{T('Baisse', 'Drawdown')} · {n}", 18)
    for i, (d, row) in enumerate(weekly.iterrows()):
        wd.cell(row=2 + i, column=1, value=d.date()).number_format = "dd/mm/yyyy"
        for j, n in enumerate(names):
            wd.cell(row=2 + i, column=2 + j, value=float(row[n])).number_format = "#,##0"
            wd.cell(row=2 + i, column=2 + len(names) + j, value=float(wdd[n].iloc[i])).number_format = "0.0%"
    last = 1 + len(weekly)
    wd.freeze_panes = "B2"
    cats = Reference(wd, min_col=1, min_row=2, max_row=last)

    growth = LineChart()
    _style_chart(growth, T(f"Évolution de {capital:,.0f}".replace(",", " "), f"Growth of {capital:,.0f}"), "#,##0", 9, 24)
    for j in range(len(names)):
        growth.add_data(Reference(wd, min_col=2 + j, min_row=1, max_row=last), titles_from_data=True)
    growth.set_categories(cats)
    growth.x_axis.number_format = "mm/yyyy"
    for j, s in enumerate(growth.series):
        s.graphicalProperties.line.solidFill = SERIES[j % len(SERIES)] if j < len(names) - 1 else "9AA29D"
        s.graphicalProperties.line.width = 22000 if j < len(names) - 1 else 15000
        if j == len(names) - 1:
            s.graphicalProperties.line.dashStyle = "dash"
        s.smooth = False
    ws.add_chart(growth, f"{get_column_letter(len(names) + 3)}5")

    ddc = LineChart()
    _style_chart(ddc, T("Baisses depuis le dernier sommet", "Drawdowns from the last peak"), "0%", 7.5, 24)
    for j in range(len(names) - 1):
        ddc.add_data(Reference(wd, min_col=2 + len(names) + j, min_row=1, max_row=last), titles_from_data=True)
    ddc.set_categories(cats)
    ddc.x_axis.number_format = "mm/yyyy"
    for j, s in enumerate(ddc.series):
        s.graphicalProperties.line.solidFill = SERIES[j % len(SERIES)]
        s.graphicalProperties.line.width = 15000
    ws.add_chart(ddc, f"{get_column_letter(len(names) + 3)}24")

    # 4) Rendements annuels + graphique en barres
    wy = wb.create_sheet(T("Rendements annuels", "Annual returns"))
    _head(wy, 1, 1, T("Année", "Year"), 10)
    for j, n in enumerate(names):
        _head(wy, 1, 2 + j, n, 16)
    for i, (y, row) in enumerate(yearly.iterrows()):
        _cell(wy, 2 + i, 1, int(y), bold=True)
        for j, n in enumerate(names):
            _cell(wy, 2 + i, 2 + j, float(row[n]), "0.0%")
    ylast = 1 + len(yearly)
    rng = f"B2:{get_column_letter(1 + len(names))}{ylast}"
    wy.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=-0.3, start_color="F8696B", mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="num", end_value=0.3, end_color="63BE7B"))
    note = ylast + 2
    wy.cell(row=note, column=1, value=T("La première année est partielle (à partir de la date de départ commune), comme la dernière.",
                                        "The first year is partial (from the shared start date), as is the last one.")).font = Font(size=9, italic=True, color="69716C")
    bars = BarChart()
    bars.type = "col"
    _style_chart(bars, T("Rendement par année civile", "Return per calendar year"), "0%", 9, 24)
    for j in range(len(names)):
        bars.add_data(Reference(wy, min_col=2 + j, min_row=1, max_row=ylast), titles_from_data=True)
    bars.set_categories(Reference(wy, min_col=1, min_row=2, max_row=ylast))
    for j, s in enumerate(bars.series):
        s.graphicalProperties.solidFill = SERIES[j % len(SERIES)] if j < len(names) - 1 else "BFC4C0"
        s.graphicalProperties.line.noFill = True
    bars.gapWidth = 60
    ws.add_chart(bars, f"A{r0 + len(rows) + 3}")
    wy.add_chart(_clone_bars(wy, names, ylast), f"{get_column_letter(len(names) + 3)}2")

    # 5) Risque / rendement (nuage de points)
    wr = wb.create_sheet(T("Risque-rendement", "Risk-return"))
    _head(wr, 1, 1, T("Portefeuille", "Portfolio"), 24)
    _head(wr, 1, 2, T("Volatilité", "Volatility"), 12)
    _head(wr, 1, 3, "CAGR", 12)
    _head(wr, 1, 4, T("Pire baisse", "Worst drop"), 12)
    for i, n in enumerate(names):
        _cell(wr, 2 + i, 1, n, bold=True)
        _cell(wr, 2 + i, 2, metrics[n]["annualized_volatility_pct"] / 100, "0.0%")
        _cell(wr, 2 + i, 3, metrics[n]["cagr_pct"] / 100, "0.0%")
        _cell(wr, 2 + i, 4, metrics[n]["max_drawdown_pct"] / 100, "0.0%")
    sc = ScatterChart()
    _style_chart(sc, T("Risque (volatilité) contre rendement annuel", "Risk (volatility) versus annual return"), "0%", 9, 16)
    sc.style = 13
    sc.x_axis.title = T("Volatilité annualisée", "Annualised volatility")
    sc.y_axis.title = "CAGR"
    sc.x_axis.numFmt = "0%"
    for i, n in enumerate(names):
        s = Series(Reference(wr, min_col=3, min_row=2 + i), Reference(wr, min_col=2, min_row=2 + i), title=n)
        s.marker.symbol = "circle"
        s.marker.size = 11
        s.marker.graphicalProperties.solidFill = SERIES[i % len(SERIES)] if i < len(names) - 1 else "9AA29D"
        s.marker.graphicalProperties.line.solidFill = s.marker.graphicalProperties.solidFill
        s.graphicalProperties.line.noFill = True
        sc.series.append(s)
    wr.add_chart(sc, "F2")
    wr.cell(row=len(names) + 3, column=1, value=T("En haut à gauche : plus de rendement pour moins de risque.", "Top-left: more return for less risk.")).font = Font(size=9, italic=True, color="69716C")

    # 6) Rendements mensuels (carte de chaleur par portefeuille)
    wm = wb.create_sheet(T("Rendements mensuels", "Monthly returns"))
    months = ["Janv", "Févr", "Mars", "Avr", "Mai", "Juin", "Juil", "Août", "Sept", "Oct", "Nov", "Déc"] if lang() == "fr" else \
        ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    r = 1
    wm.column_dimensions["A"].width = 12
    for n in port_names:
        wm.cell(row=r, column=1, value=n).font = Font(bold=True, size=13, color=GREEN)
        r += 1
        _head(wm, r, 1, T("Année", "Year"))
        for k, mname in enumerate(months):
            _head(wm, r, 2 + k, mname, 8)
        _head(wm, r, 14, T("Année", "Year"), 9)
        first = r + 1
        m = monthly[n]
        for y in sorted(set(m.index.year)):
            r += 1
            _cell(wm, r, 1, int(y), bold=True)
            for k in range(12):
                v = m[(m.index.year == y) & (m.index.month == k + 1)]
                _cell(wm, r, 2 + k, float(v.iloc[0]) if len(v) else None, "0.0%")
            yv = yearly[n].get(y)
            _cell(wm, r, 14, float(yv) if yv is not None else None, "0.0%", bold=True)
        wm.conditional_formatting.add(f"B{first}:M{r}", ColorScaleRule(start_type="num", start_value=-0.1, start_color="F8696B", mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="num", end_value=0.1, end_color="63BE7B"))
        r += 3

    # 7) Compositions + camemberts
    wc = wb.create_sheet(T("Compositions", "Compositions"))
    wc.column_dimensions["A"].width = 14
    wc.column_dimensions["B"].width = 34
    wc.column_dimensions["C"].width = 10
    r = 1
    for idx, p in enumerate(portfolios):
        wc.cell(row=r, column=1, value=p["name"]).font = Font(bold=True, size=13, color=GREEN)
        wc.cell(row=r, column=2, value=f"{T('Rééquilibrage', 'Rebalancing')} : {p.get('rebalance') or 'none'}").font = Font(size=9, color="69716C")
        r += 1
        _head(wc, r, 1, "Ticker"); _head(wc, r, 2, T("Actif", "Asset")); _head(wc, r, 3, T("Poids", "Weight"))
        first = r + 1
        for tk, wgt in p["items"]:
            r += 1
            _cell(wc, r, 1, tk, bold=True); _cell(wc, r, 2, asset_name(tk)); _cell(wc, r, 3, float(wgt) / 100, "0.0%")
        pie = PieChart()
        pie.title = p["name"]
        pie.height, pie.width = 6.5, 9
        pie.add_data(Reference(wc, min_col=3, min_row=first, max_row=r))
        pie.set_categories(Reference(wc, min_col=1, min_row=first, max_row=r))
        pie.dataLabels = DataLabelList()
        pie.dataLabels.showPercent = True
        wc.add_chart(pie, f"E{first - 1}")
        r = max(r, first + 12) + 2

    # 8) Corrélations entre portefeuilles
    wk = wb.create_sheet(T("Corrélations", "Correlations"))
    corr = rets.corr()
    _head(wk, 1, 1, "", 24)
    for j, n in enumerate(names):
        _head(wk, 1, 2 + j, n, 16)
    for i, a in enumerate(names):
        _cell(wk, 2 + i, 1, a, bold=True)
        for j, b in enumerate(names):
            _cell(wk, 2 + i, 2 + j, float(corr.loc[a, b]), "0.00")
    wk.conditional_formatting.add(f"B2:{get_column_letter(1 + len(names))}{1 + len(names)}",
                                  ColorScaleRule(start_type="num", start_value=-1, start_color="5B9BD5", mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="num", end_value=1, end_color="F8696B"))
    wk.cell(row=len(names) + 3, column=1, value=T("Proche de 1 : les portefeuilles évoluent ensemble, les combiner diversifie peu.",
                                                   "Close to 1: the portfolios move together, combining them adds little diversification.")).font = Font(size=9, italic=True, color="69716C")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _clone_bars(ws, names, ylast):
    """Même graphique en barres, ancré dans la feuille des rendements annuels."""
    bars = BarChart()
    bars.type = "col"
    _style_chart(bars, T("Rendement par année civile", "Return per calendar year"), "0%", 9, 20)
    for j in range(len(names)):
        bars.add_data(Reference(ws, min_col=2 + j, min_row=1, max_row=ylast), titles_from_data=True)
    bars.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ylast))
    for j, s in enumerate(bars.series):
        s.graphicalProperties.solidFill = SERIES[j % len(SERIES)] if j < len(names) - 1 else "BFC4C0"
        s.graphicalProperties.line.noFill = True
    bars.gapWidth = 60
    return bars
